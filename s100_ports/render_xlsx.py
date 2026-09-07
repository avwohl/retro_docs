# -*- coding: utf-8 -*-
"""Render the port dataset as an Excel workbook.

Run `python build.py` to regenerate both deliverables from ports_data.py.
"""
import pathlib, datetime
OUT_DIR = pathlib.Path(__file__).resolve().parent

from ports_data import MITS, THIRD, MEMMAP
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INK      = "1F2937"
HDR_BG   = "1F3A5F"
HDR_FG   = "FFFFFF"
BAND     = "F2F5F9"
WARN_BG  = "FDECEC"
NOTE_BG  = "FFF7E0"
RULE     = Side(style="thin", color="D3D9E2")

CAT_FILL = {
    "Serial": "E3F0FF", "Printer": "FFEFE0", "Floppy": "E6F6E9", "Hard disk": "DCEEDF",
    "Parallel": "F1E8FB", "Video": "FFE9F3", "Cassette": "FFF6D9", "Keyboard": "FFF6D9",
    "Interrupt / clock": "EDEDED", "System": "EDEDED", "Timer": "EDEDED", "Clock": "EDEDED",
    "Math": "EDEDED", "DMA": "E8E8F5", "Modem": "E0F5F3", "Analog I/O": "F1E8FB",
    "RAM disk": "DCEEDF", "Serial / Parallel": "E3F0FF", "Floppy / Hard disk": "E6F6E9",
}

def octal(n):  return format(n, "03o")

def fmt(port, base):
    """Return (hex, dec, oct) display strings for an int, (lo,hi) tuple, or None."""
    if port is None:
        return ("-", "-", "-")
    if isinstance(port, tuple):
        lo, hi = port
        return ("0x%02X-0x%02X" % (lo, hi), "%d-%d" % (lo, hi), "%s-%s" % (octal(lo), octal(hi)))
    return ("0x%02X" % port, str(port), octal(port))

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=HDR_FG, size=10)
        cell.fill = PatternFill("solid", fgColor=HDR_BG)
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=HDR_BG))
    ws.row_dimensions[row].height = 30
    # NB: use a string ref - ws.cell() would advance _current_row and leave a blank row
    ws.freeze_panes = "A%d" % (row + 1)

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

wb = Workbook()

# ---------------------------------------------------------------- About ----
ws = wb.active
ws.title = "About"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 3, "B": 26, "C": 104})

def about_row(r, label, text, bold=False, fill=None, height=None):
    lc, tc = ws.cell(row=r, column=2), ws.cell(row=r, column=3)
    lc.value, tc.value = label, text
    lc.font = Font(bold=True, size=10, color=INK)
    tc.font = Font(bold=bold, size=10, color=INK)
    lc.alignment = Alignment(vertical="top")
    tc.alignment = Alignment(vertical="top", wrap_text=True)
    if fill:
        for c in (lc, tc):
            c.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[r].height = height

t = ws.cell(row=2, column=2, value="Altair / S-100 Bus I/O Port Assignments")
t.font = Font(bold=True, size=18, color=HDR_BG)
s = ws.cell(row=3, column=2, value="MITS Altair factory hardware and third-party S-100 cards - serial, printer, floppy, hard disk and support boards")
s.font = Font(size=11, italic=True, color="5A6472")
ws.cell(row=4, column=2, value="Compiled " + datetime.date.today().isoformat()).font = Font(size=9, color="8A94A2")

r = 6
about_row(r, "Sheets", "MITS Altair - every MITS/Altair board with a documented port.  Third-Party Cards - the major S-100 makers.  "
                       "Port Map 00-FF - a byte-by-byte index showing who claims each port and where they collide.  "
                       "Memory-Mapped - controllers that live in memory space and use no ports at all.  Sources - what each row is based on.", height=60); r += 2

about_row(r, "Read this first", "Almost nothing on the S-100 bus had a fixed address. Boards decoded their port range with jumpers or a DIP "
                                "switch, so the addresses here are documented factory defaults and de-facto standards - the settings software "
                                "expected to find - not hard-wired facts. Any given machine may differ. Where a board's address is genuinely "
                                "fixed the Addressing column says so.", height=62); r += 2

about_row(r, "Octal matters", "MITS documented ports in OCTAL, and so did most Altair-era software listings. The 88-2SIO console at 0x10/0x11 "
                              "is 'port 20 and 21' in every MITS manual. All three radixes are given on every row so you can read a period "
                              "listing without converting in your head.", height=48); r += 2

about_row(r, "The core four", "For the systems this document is really about, the addresses worth memorising are: console serial 0x10/0x11 "
                             "(88-2SIO) or 0x00/0x01 (88-SIO); printer 0x02/0x03; floppy 0x08/0x09/0x0A (88-DCDD) or 0xF8-0xFC (Tarbell) "
                             "or 0x30-0x34 (Cromemco); hard disk 0xA0-0xA7 (Altair 88-HDSK via a 4PIO).", height=48); r += 2

about_row(r, "Known collisions", "The low ports were badly overcrowded. The 88-4PIO's own manual example (0x10-0x1F) sits on top of the standard "
                                 "88-2SIO console. The Cromemco disk controller's 0x03/0x04 collide with IMSAI's SIO-2 console convention. "
                                 "0xFD is claimed by both the IMSAI FIF and the Tarbell 2022. The Port Map sheet flags every one of these in red.",
          fill=WARN_BG, height=50); r += 2

about_row(r, "Confidence column", "High = stated in a manufacturer manual, or agreed on by two independent emulator implementations. "
                                  "Medium-High = one strong source. Medium = one source, or a documented conflict between sources. "
                                  "Unverified = the board is listed for completeness but no address was confirmed; these rows are deliberately "
                                  "left blank rather than guessed at.", height=50); r += 2

about_row(r, "Coverage", "This covers the manufacturers whose cards you actually meet in Altair-era systems. Several hundred S-100 boards were "
                         "made in total, and a long tail of them - mostly memory, prototyping and one-off cards - is not represented here.",
          fill=NOTE_BG, height=38)

# ------------------------------------------------------- Board sheets ------
BOARD_HDR = ["Manufacturer", "Board / Model", "Category", "Port (hex)", "Dec", "Octal",
             "R/W", "Function", "Notes", "Addressing", "Confidence", "Source"]

def board_sheet(title, rows):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.append(BOARD_HDR)
    style_header(ws, len(BOARD_HDR))
    prev_board = None
    band = False
    for (mfr, board, cat, port, dirn, func, notes, addressing, conf, src) in rows:
        if board != prev_board:
            band = not band
            prev_board = board
        h, d, o = fmt(port, None)
        ws.append([mfr, board, cat, h, d, o, dirn, func, notes, addressing, conf, src])
        ridx = ws.max_row
        for c in range(1, len(BOARD_HDR) + 1):
            cell = ws.cell(row=ridx, column=c)
            cell.font = Font(size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (2, 8, 9, 10, 12)))
            cell.border = Border(bottom=RULE)
            if band:
                cell.fill = PatternFill("solid", fgColor=BAND)
        ws.cell(row=ridx, column=3).fill = PatternFill("solid", fgColor=CAT_FILL.get(cat, "EDEDED"))
        for c in (4, 5, 6, 7):
            ws.cell(row=ridx, column=c).font = Font(size=10, color=INK, name="Menlo")
            ws.cell(row=ridx, column=c).alignment = Alignment(vertical="top", horizontal="left")
        if conf == "Unverified":
            for c in range(1, len(BOARD_HDR) + 1):
                ws.cell(row=ridx, column=c).fill = PatternFill("solid", fgColor=NOTE_BG)
        ws.cell(row=ridx, column=11).font = Font(size=10, color=INK,
                                                bold=(conf in ("Unverified",)))
    widths(ws, {"A": 20, "B": 30, "C": 15, "D": 12, "E": 9, "F": 10, "G": 6,
                "H": 62, "I": 52, "J": 26, "K": 13, "L": 40})
    ws.auto_filter.ref = "A1:L%d" % ws.max_row
    return ws

board_sheet("MITS Altair", MITS)
board_sheet("Third-Party Cards", THIRD)

# ----------------------------------------------------------- Port map ------
ws = wb.create_sheet("Port Map 00-FF")
ws.sheet_view.showGridLines = False
ws.append(["Port (hex)", "Dec", "Octal", "Claimed by", "Count", "Category"])
style_header(ws, 6)

claims = {}
for src_rows in (MITS, THIRD):
    for (mfr, board, cat, port, dirn, func, notes, addressing, conf, s) in src_rows:
        if port is None:
            continue
        ports = range(port[0], port[1] + 1) if isinstance(port, tuple) else [port]
        for p in ports:
            label = "%s %s" % (mfr, board)
            claims.setdefault(p, [])
            if label not in [c[0] for c in claims[p]]:
                claims[p].append((label, cat))

for p in range(256):
    entries = claims.get(p, [])
    names = sorted({n for n, _ in entries})
    cats = sorted({c for _, c in entries})
    ws.append(["0x%02X" % p, p, octal(p), "  |  ".join(names) if names else "",
               len(names), ", ".join(cats)])
    ridx = ws.max_row
    for c in range(1, 7):
        cell = ws.cell(row=ridx, column=c)
        cell.font = Font(size=10, color=INK, name="Menlo" if c in (1, 2, 3) else None)
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (4, 6)))
        cell.border = Border(bottom=RULE)
        if len(names) > 1:
            cell.fill = PatternFill("solid", fgColor=WARN_BG)
        elif not names:
            cell.fill = PatternFill("solid", fgColor="FAFBFC")
    if len(names) > 1:
        ws.cell(row=ridx, column=5).font = Font(size=10, bold=True, color="B42318")

widths(ws, {"A": 12, "B": 8, "C": 9, "D": 96, "E": 8, "F": 26})
ws.auto_filter.ref = "A1:F%d" % ws.max_row

# ------------------------------------------------------- Memory-mapped -----
ws = wb.create_sheet("Memory-Mapped")
ws.sheet_view.showGridLines = False
ws.append(["Manufacturer", "Board", "Type", "Memory range", "Function", "Notes", "Source"])
style_header(ws, 7)
for (mfr, board, cat, rng, func, notes, src) in MEMMAP:
    ws.append([mfr, board, cat, rng, func, notes, src])
    ridx = ws.max_row
    for c in range(1, 8):
        cell = ws.cell(row=ridx, column=c)
        cell.font = Font(size=10, color=INK, name="Menlo" if c == 4 else None)
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (2, 5, 6, 7)))
        cell.border = Border(bottom=RULE)
widths(ws, {"A": 22, "B": 32, "C": 20, "D": 22, "E": 46, "F": 62, "G": 34})
ws.auto_filter.ref = "A1:G%d" % ws.max_row

# ------------------------------------------------------------ Sources ------
ws = wb.create_sheet("Sources")
ws.sheet_view.showGridLines = False
ws.append(["Source", "What it establishes", "URL"])
style_header(ws, 3)
SOURCES = [
 ("MITS 88-4PIO manual", "Board occupies 16 consecutive addresses; jumper groups at 000/020/040/060/100/120... octal; per-PIA register layout; worked example at 020-037 octal", "http://dunfield.classiccmp.org/s100c/mits/88_4pio.pdf"),
 ("MITS 88-VI/RTC manual", "'The 88-VI (RTC) uses I/O address 254 (decimal) or 376 (octal)'", "http://dunfield.classiccmp.org/s100c/mits/88_virtc.pdf"),
 ("MITS 88-SIO manual", "Two device addresses, jumper-selectable to any even octal address 000-376; A0 selects control vs data", "http://dunfield.classiccmp.org/s100c/mits/88sio_2.pdf"),
 ("MITS 88-ACR manual", "Cassette interface addressing and octal selection charts", "http://dunfield.classiccmp.org/s100c/mits/88_acr.pdf"),
 ("SIMH AltairZ80 documentation", "88-2SIO at 10-13 hex as the MITS standard; 88-DISK at 8/9/0A; North Star and Micropolis memory windows; CompuPro and Cromemco board descriptions", "http://bitsavers.trailing-edge.com/simh.trailing-edge.com_201206/pdf/altairz80_doc.pdf"),
 ("SIMH AltairZ80 device sources", "Register-level maps for Tarbell, Cromemco 4/16/64FDC, CompuPro Disk1A/2/3/Selector/M-Drive/System Support 1, Interfacer 3, Morrow, iCOM, Jade, PMMI, Hayes, Sol-20, TU-ART, Dazzler", "https://github.com/simh/simh/tree/master/AltairZ80"),
 ("Altair 8800 Simulator (D. Hansel)", "Independent confirmation of the MITS port map (SIO 0/1, printer 2/3, ACR 6/7, DCDD 8/9/0A, 2SIO 10-13, VI 0xFE, sense switches 0xFF), 88-HDSK via a 4PIO at 0xA0-0xA7, Tarbell F8-FD, Cromemco 30-34", "https://github.com/dhansel/Altair8800"),
 ("Altair 8800 Simulator manual", "88-SIO at 0/1, 88-ACR at 6/7, 88-2SIO at 10/11 and 12/13; printer control port 02h", "https://retrocmp.de/hardware/altair-8800/altair-8800-sim-manual.pdf"),
 ("Cromemco Cromix Instruction Manual (023-4022)", "TU-ART switch settings and the multi-user port map: #1 A=20h B=50h, #2 A=60h B=70h, #3 A=80h", "https://archive.org/details/023-4022-cromemco-cromix-manuals"),
 ("CompuPro Interfacer 4 Technical Manual (187C)", "Eight-port block on any multiple of 8; CompuPro default 10-17h; relative port 0-7 function table", "https://archive.org/details/bitsavers_compupro18calManualMay83_3167199"),
 ("IMSAI CP/M System User's Guide and SIO-2 manual", "SIO-2 default block 00-0F; IMSAI software uses 02/03 (TTY) and 04/05 (CRT); second board at 20-2F", "http://www.bitsavers.org/pdf/imsai/IMSAI_SIO2-2_B_Manual.pdf"),
 ("s100computers.com", "Board histories and addressing notes for MITS, Tarbell, Cromemco, CompuPro and Processor Technology cards", "http://www.s100computers.com/"),
 ("deramp.com archive", "North Star Horizon restoration notes (MDS controller occupies E800-EFFF); MITS and IMSAI documentation mirror", "https://deramp.com/"),
]
for (name, what, url) in SOURCES:
    ws.append([name, what, url])
    ridx = ws.max_row
    for c in range(1, 4):
        cell = ws.cell(row=ridx, column=c)
        cell.font = Font(size=10, color=INK)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=RULE)
    u = ws.cell(row=ridx, column=3)
    u.hyperlink = url
    u.font = Font(size=10, color="1155CC", underline="single")
widths(ws, {"A": 38, "B": 78, "C": 58})

out = str(OUT_DIR / "Altair_S100_Port_Assignments.xlsx")
wb.save(out)
print("wrote", out)
print("MITS rows:", len(MITS), "| third-party rows:", len(THIRD), "| memory-mapped:", len(MEMMAP))
conflicts = sorted(p for p, v in claims.items() if len({n for n, _ in v}) > 1)
print("ports with >1 claimant:", ", ".join("0x%02X" % p for p in conflicts))
print("ports claimed at all:", len(claims))
